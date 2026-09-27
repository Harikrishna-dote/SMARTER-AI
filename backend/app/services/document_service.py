import asyncio
import re
from pathlib import Path
from typing import Any
from uuid import uuid4

import logging
from docx import Document as WordDocument
from fastapi import HTTPException, UploadFile, status
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.document import Document
from app.repositories.documents import DocumentRepository
from app.services.ai_gateway import AIGateway
from app.services.memory_service import MemoryService


class DocumentService:
    _ALLOWED_SUFFIXES = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".txt": "text/plain",
        ".md": "text/markdown",
    }

    def __init__(self, db: AsyncSession, gateway: AIGateway | None = None, vector_service: Any | None = None):
        self.repo = DocumentRepository(db)
        self.gateway = gateway or AIGateway()
        self.vector_service = vector_service
        self.settings = get_settings()

    async def save_upload(self, user_id: str, file: UploadFile) -> Document:
        upload_dir = Path(self.settings.upload_dir)
        upload_dir.mkdir(parents=True, exist_ok=True)
        filename = self._sanitize_filename(file.filename)
        suffix = Path(filename).suffix.lower()
        content_type = self._validate_upload(filename, file.content_type)
        safe_name = f"{uuid4()}{suffix}"
        path = upload_dir / safe_name
        try:
            await self._write_upload(path, file)
            extracted_text = await asyncio.to_thread(self.extract_text, path, content_type)
            document = Document(
                user_id=user_id,
                filename=filename,
                content_type=content_type,
                storage_path=str(path),
                extracted_text=extracted_text,
            )
            saved = await self.repo.add(document)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        await MemoryService(self.repo.db).persist_document_memory(
            user_id,
            saved.id,
            saved.filename,
            saved.extracted_text,
        )
        if self.vector_service and saved.extracted_text:
            await self._index_document(saved)
        return saved

    def _sanitize_filename(self, filename: str | None) -> str:
        raw_name = Path(filename or "upload").name
        safe_name = re.sub(r"[\x00-\x1f\x7f]+", "", raw_name).strip()
        if not safe_name or safe_name in {".", ".."}:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="A valid filename is required")
        return safe_name[:255]

    def _validate_upload(self, filename: str, content_type: str | None) -> str:
        normalized_type = (content_type or "").split(";", 1)[0].strip().lower()
        expected_type = self._ALLOWED_SUFFIXES.get(Path(filename).suffix.lower())
        allowed_types = {item.lower() for item in self.settings.allowed_upload_types}
        if expected_type is None or expected_type not in allowed_types:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="Unsupported file type")
        if normalized_type and normalized_type != expected_type:
            raise HTTPException(status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, detail="File type does not match its extension")
        return expected_type

    async def _write_upload(self, path: Path, file: UploadFile) -> None:
        max_bytes = max(1, int(self.settings.max_upload_size_mb)) * 1024 * 1024
        written = 0
        with path.open("wb") as destination:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                written += len(chunk)
                if written > max_bytes:
                    raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="File is too large")
                destination.write(chunk)

    def extract_text(self, path: Path, content_type: str) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf" or content_type == "application/pdf":
            reader = PdfReader(str(path))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if suffix in {".docx", ".doc"}:
            document = WordDocument(str(path))
            return "\n".join(paragraph.text for paragraph in document.paragraphs)
        if suffix in {".txt", ".md", ".csv", ".json"}:
            return path.read_text(encoding="utf-8", errors="ignore")
        return ""

    async def answer_question(self, user_id: str, document_id: str, question: str) -> str:
        document = await self.repo.get_for_user(document_id, user_id)
        if not document:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")
        context = document.extracted_text[:16000]
        return await self.gateway.chat(
            [
                {"role": "system", "content": "Answer using the supplied document context. Say when the answer is not present."},
                {"role": "user", "content": f"Document:\n{context}\n\nQuestion: {question}"},
            ]
        )

    async def _index_document(self, document: Document) -> None:
        if not self.vector_service:
            return
        try:
            text = document.extracted_text
            chunks = self._chunk_text(text, chunk_size=1000, overlap=200)
            for chunk_index, chunk in enumerate(chunks):
                await self.vector_service.index_document(
                    document_id=document.id,
                    text=chunk,
                    metadata={
                        "document_id": document.id,
                        "chunk_index": chunk_index,
                        "chunk_count": len(chunks),
                        "filename": document.filename,
                        "content_type": document.content_type,
                        "user_id": document.user_id,
                        "source": document.filename,
                    },
                )
        except Exception as exc:
            logger = logging.getLogger(__name__)
            logger.warning("Failed to index document %s: %s", document.id, exc)

    @staticmethod
    def _chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
        if not text:
            return []
        chunks: list[str] = []
        start = 0
        while start < len(text):
            end = start + chunk_size
            chunk = text[start:end]
            if chunk.strip():
                chunks.append(chunk)
            start = end - overlap
        return chunks
