"""Context-aware translation, grammar, and language analysis."""

import asyncio
import json
import logging
import random
import re
import time
from datetime import datetime
from enum import Enum
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
from fastapi import HTTPException, UploadFile, status
from PIL import Image, UnidentifiedImageError
from pydantic import BaseModel, ValidationError
from pypdf import PdfReader
from docx import Document as WordDocument
from pytesseract import TesseractNotFoundError
from sqlalchemy import Column, DateTime, Integer, String, Text, desc, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.cache import fast_response_cache, translation_cache
from app.core.database import Base
from app.models.base import new_id
from app.models.translation import (
    Bookmark,
    Flashcard,
    OCRElement,
    Quiz,
    SavedPhrase,
    TranslationHistory,
    VocabularyItem,
)
from app.schemas.translation import (
    AnalysisRequest,
    AnalysisResponse,
    BookmarkResponse,
    ExplainRequest,
    ExplainResponse,
    FlashcardGenerateRequest,
    FlashcardResponse,
    GrammarRequest,
    GrammarResponse,
    LanguageOption,
    OCRBox,
    OCRResponse,
    QuizGenerateRequest,
    QuizResponse,
    SavedPhraseResponse,
    SentenceBreakdown,
    SentenceAnalysis,
    TranslationTextRequest,
    TranslationTextResponse,
    VocabularyItemResponse,
    WordAnalysis,
)
from app.services.ai_gateway import AIGateway
from app.services.vision_service import VisionService


logger = logging.getLogger(__name__)


LANGUAGES: tuple[tuple[str, str, str], ...] = (
    ("af", "Afrikaans", "Afrikaans"), ("sq", "Albanian", "Shqip"),
    ("am", "Amharic", "አማርኛ"), ("ar", "Arabic", "العربية"),
    ("hy", "Armenian", "Հայերեն"), ("as", "Assamese", "অসমীয়া"),
    ("ay", "Aymara", "Aymar aru"), ("az", "Azerbaijani", "Azərbaycan"),
    ("bm", "Bambara", "Bamanankan"), ("eu", "Basque", "Euskara"),
    ("be", "Belarusian", "Беларуская"), ("bn", "Bengali", "বাংলা"),
    ("bho", "Bhojpuri", "भोजपुरी"), ("bs", "Bosnian", "Bosanski"),
    ("bg", "Bulgarian", "Български"), ("ca", "Catalan", "Català"),
    ("ceb", "Cebuano", "Cebuano"), ("zh", "Chinese (Simplified)", "简体中文"),
    ("zh-TW", "Chinese (Traditional)", "繁體中文"), ("co", "Corsican", "Corsu"),
    ("hr", "Croatian", "Hrvatski"), ("cs", "Czech", "Čeština"),
    ("da", "Danish", "Dansk"), ("dv", "Dhivehi", "ދިވެހި"),
    ("doi", "Dogri", "डोगरी"), ("nl", "Dutch", "Nederlands"),
    ("en", "English", "English"), ("eo", "Esperanto", "Esperanto"),
    ("et", "Estonian", "Eesti"), ("ee", "Ewe", "Eʋegbe"),
    ("fil", "Filipino", "Filipino"), ("fi", "Finnish", "Suomi"),
    ("fr", "French", "Français"), ("fy", "Frisian", "Frysk"),
    ("gl", "Galician", "Galego"), ("ka", "Georgian", "ქართული"),
    ("de", "German", "Deutsch"), ("el", "Greek", "Ελληνικά"),
    ("gn", "Guarani", "Avañe'ẽ"), ("gu", "Gujarati", "ગુજરાતી"),
    ("ht", "Haitian Creole", "Kreyòl ayisyen"), ("ha", "Hausa", "Hausa"),
    ("haw", "Hawaiian", "ʻŌlelo Hawaiʻi"), ("he", "Hebrew", "עברית"),
    ("hi", "Hindi", "हिन्दी"), ("hmn", "Hmong", "Hmoob"),
    ("hu", "Hungarian", "Magyar"), ("is", "Icelandic", "Íslenska"),
    ("ig", "Igbo", "Igbo"), ("ilo", "Ilocano", "Ilocano"),
    ("id", "Indonesian", "Bahasa Indonesia"), ("ga", "Irish", "Gaeilge"),
    ("it", "Italian", "Italiano"), ("ja", "Japanese", "日本語"),
    ("jv", "Javanese", "Basa Jawa"), ("kn", "Kannada", "ಕನ್ನಡ"),
    ("kk", "Kazakh", "Қазақша"), ("km", "Khmer", "ខ្មែរ"),
    ("rw", "Kinyarwanda", "Ikinyarwanda"), ("gom", "Konkani", "कोंकणी"),
    ("ko", "Korean", "한국어"), ("kri", "Krio", "Krio"),
    ("ku", "Kurdish", "Kurdî"), ("ckb", "Kurdish (Sorani)", "کوردی"),
    ("ky", "Kyrgyz", "Кыргызча"), ("lo", "Lao", "ລາວ"),
    ("la", "Latin", "Latina"), ("lv", "Latvian", "Latviešu"),
    ("ln", "Lingala", "Lingála"), ("lt", "Lithuanian", "Lietuvių"),
    ("lg", "Luganda", "Luganda"), ("lb", "Luxembourgish", "Lëtzebuergesch"),
    ("mk", "Macedonian", "Македонски"), ("mai", "Maithili", "मैथिली"),
    ("mg", "Malagasy", "Malagasy"), ("ms", "Malay", "Bahasa Melayu"),
    ("ml", "Malayalam", "മലയാളം"), ("mt", "Maltese", "Malti"),
    ("mi", "Maori", "Māori"), ("mr", "Marathi", "मराठी"),
    ("mni-Mtei", "Meiteilon", "ꯃꯤꯇꯩꯂꯣꯟ"), ("lus", "Mizo", "Mizo ṭawng"),
    ("mn", "Mongolian", "Монгол"), ("my", "Myanmar", "မြန်မာ"),
    ("ne", "Nepali", "नेपाली"), ("no", "Norwegian", "Norsk"),
    ("ny", "Nyanja", "Chichewa"), ("or", "Odia", "ଓଡ଼ିଆ"),
    ("om", "Oromo", "Afaan Oromoo"), ("ps", "Pashto", "پښتو"),
    ("fa", "Persian", "فارسی"), ("pl", "Polish", "Polski"),
    ("pt", "Portuguese", "Português"), ("pa", "Punjabi", "ਪੰਜਾਬੀ"),
    ("qu", "Quechua", "Runasimi"), ("ro", "Romanian", "Română"),
    ("ru", "Russian", "Русский"), ("sm", "Samoan", "Gagana Samoa"),
    ("sa", "Sanskrit", "संस्कृतम्"), ("gd", "Scots Gaelic", "Gàidhlig"),
    ("nso", "Sepedi", "Sepedi"), ("sr", "Serbian", "Српски"),
    ("st", "Sesotho", "Sesotho"), ("sn", "Shona", "ChiShona"),
    ("sd", "Sindhi", "سنڌي"), ("si", "Sinhala", "සිංහල"),
    ("sk", "Slovak", "Slovenčina"), ("sl", "Slovenian", "Slovenščina"),
    ("so", "Somali", "Soomaali"), ("es", "Spanish", "Español"),
    ("su", "Sundanese", "Basa Sunda"), ("sw", "Swahili", "Kiswahili"),
    ("sv", "Swedish", "Svenska"), ("tl", "Tagalog", "Tagalog"),
    ("tg", "Tajik", "Тоҷикӣ"), ("ta", "Tamil", "தமிழ்"),
    ("tt", "Tatar", "Татарча"), ("te", "Telugu", "తెలుగు"),
    ("th", "Thai", "ไทย"), ("ti", "Tigrinya", "ትግርኛ"),
    ("ts", "Tsonga", "itsonga"), ("tr", "Turkish", "Türkçe"),
    ("tk", "Turkmen", "Türkmençe"), ("ak", "Twi", "Twi"),
    ("uk", "Ukrainian", "Українська"), ("ur", "Urdu", "اردو"),
    ("ug", "Uyghur", "ئۇيغۇرچە"), ("uz", "Uzbek", "O‘zbek"),
    ("vi", "Vietnamese", "Tiếng Việt"), ("cy", "Welsh", "Cymraeg"),
    ("xh", "Xhosa", "isiXhosa"), ("yi", "Yiddish", "ייִדיש"),
    ("yo", "Yoruba", "Yorùbá"), ("zu", "Zulu", "isiZulu"),
)

LANGUAGE_BY_CODE = {code.lower(): (code, name, native) for code, name, native in LANGUAGES}
LANGUAGE_BY_NAME = {
    label.lower(): code
    for code, name, native in LANGUAGES
    for label in (name, native)
}


class LanguageCode(str, Enum):
    AUTO = "auto"
    EN = "en"
    ES = "es"
    FR = "fr"
    DE = "de"
    HI = "hi"
    TE = "te"
    TA = "ta"
    KN = "kn"
    AR = "ar"
    ZH = "zh"
    JA = "ja"
    KO = "ko"


class TranslationMode(str, Enum):
    TEXT = "text"
    VOICE = "voice"
    DOCUMENT = "document"
    IMAGE = "image"
    OCR = "ocr"
    LIVE_CONVERSATION = "live_conversation"


class ContentType(str, Enum):
    PLAIN_TEXT = "plain_text"
    MARKDOWN = "markdown"
    HTML = "html"
    PDF = "pdf"
    DOCX = "docx"
    AUDIO_MP3 = "audio/mp3"
    AUDIO_WAV = "audio/wav"
    IMAGE_JPG = "image/jpg"
    IMAGE_PNG = "image/png"


class ToneStyle(str, Enum):
    PRESERVE = "preserve"
    NATURAL = "natural"
    FORMAL = "formal"
    INFORMAL = "casual"
    CONVERSATIONAL = "friendly"
    ACADEMIC = "academic"
    TECHNICAL = "technical"
    PROFESSIONAL = "professional"


class TranslationRequest(BaseModel):
    source_language: LanguageCode
    target_language: LanguageCode
    content: str
    mode: TranslationMode = TranslationMode.TEXT
    content_type: ContentType = ContentType.PLAIN_TEXT
    tone: ToneStyle | None = ToneStyle.PRESERVE
    preserve_formatting: bool = True
    context_aware: bool = True


TranslationResult = TranslationTextResponse


class VoiceTranslationResult(BaseModel):
    translated_text: str
    original_text: str
    source_language: str
    target_language: str
    audio_url: str | None = None


def supported_languages() -> list[LanguageOption]:
    return [LanguageOption(code=code, name=name, native_name=native) for code, name, native in LANGUAGES]


def _clean_json(raw: str) -> dict[str, Any]:
    cleaned = raw.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        parsed = json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", cleaned, re.DOTALL)
        if not match:
            match = re.search(r"\{.*", cleaned, re.DOTALL)
        if not match:
            raise ValueError("The translation model returned an invalid response")
        try:
            parsed = json.loads(match.group(0))
        except json.JSONDecodeError:
            raise ValueError("The translation model returned an invalid response")
    if not isinstance(parsed, dict):
        raise ValueError("The translation model returned an invalid response")
    return parsed


class TranslationEngine:
    def __init__(self, session: AsyncSession | None = None, gateway: AIGateway | None = None):
        self.session = session
        self.gateway = gateway or AIGateway()

    @staticmethod
    def validate_language(code: str, *, allow_auto: bool = False) -> str:
        normalized = code.strip()
        if allow_auto and normalized.lower() == "auto":
            return "auto"
        match = LANGUAGE_BY_CODE.get(normalized.lower())
        if not match:
            raise HTTPException(status_code=422, detail=f"Unsupported language code: {code}")
        return match[0]

    async def translate(self, request: TranslationTextRequest) -> TranslationTextResponse:
        source = self.validate_language(request.source_language, allow_auto=True)
        target = self.validate_language(request.target_language)
        start = time.perf_counter()
        chunks = self._split_content(request.content)
        payloads: list[dict[str, Any]] = []
        for index, chunk in enumerate(chunks):
            payload = None
            for chunk_attempt in range(3):
                try:
                    payload = await self._translate_chunk(
                        request,
                        chunk,
                        source,
                        target,
                        include_alternatives=request.include_alternatives and len(chunks) == 1,
                        chunk_number=index + 1,
                        chunk_count=len(chunks),
                    )
                    translation = str(payload.get("translation", "")).strip()
                    if translation:
                        break
                except HTTPException:
                    raise
                except Exception as exc:
                    logger.warning("Translation chunk %d failed (attempt %d): %s", index + 1, chunk_attempt + 1, exc)
                if chunk_attempt < 2:
                    await asyncio.sleep(min(1 + chunk_attempt, 3))
            if payload is None:
                payload = {}
            payloads.append(payload)

        translations = [str(payload.get("translation", "")).strip() for payload in payloads]
        if any(not translated for translated in translations):
            raise HTTPException(status_code=502, detail="The translation model returned an empty translation")
        first = payloads[0]
        detected = self._normalize_detected_language(str(first.get("detected_source_language", source)))
        notes = self._string_list(first.get("notes"))
        sentence_breakdown = (
            self._sentence_breakdown(first.get("sentence_breakdown"), original=request.content, translation=translations[0])
            if len(chunks) == 1
            else None
        )
        if len(chunks) > 1:
            notes.insert(0, f"Translated in {len(chunks)} sections to fit the local model context window.")
        return TranslationTextResponse(
            original_content=request.content,
            translated_content="\n\n".join(translations),
            source_language=source,
            detected_source_language=detected,
            target_language=target,
            tone=request.tone,
            transliteration=str(first.get("transliteration", "")).strip() if len(chunks) == 1 else "",
            alternatives=self._string_list(first.get("alternatives")) if len(chunks) == 1 else [],
            notes=notes,
            sentence_breakdown=sentence_breakdown,
            confidence_score=min(self._confidence(payload.get("confidence")) for payload in payloads),
            time_taken_seconds=round(time.perf_counter() - start, 3),
            character_count=len(request.content),
        )

    async def _translate_chunk(
        self,
        request: TranslationTextRequest,
        content: str,
        source: str,
        target: str,
        *,
        include_alternatives: bool,
        chunk_number: int,
        chunk_count: int,
    ) -> dict[str, Any]:
        source_label = "automatically detect the source language" if source == "auto" else self.language_name(source)
        target_label = self.language_name(target)
        formatting = (
            "Preserve paragraphs, line breaks, lists, markdown, numbers, placeholders, and punctuation."
            if request.preserve_formatting
            else "Prioritize natural target-language prose over the original layout."
        )
        context = f"\nDomain or situational context: {request.context}" if request.context.strip() else ""
        alternatives = "Provide up to 3 useful alternatives." if include_alternatives else "Return no alternatives."
        section = (
            f"\nThis is section {chunk_number} of {chunk_count}. Translate only this section and do not add section labels."
            if chunk_count > 1
            else ""
        )

        return await self._json_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are SMARTER Translate, an expert human-quality translator. Translate meaning and intent, "
                        "not words mechanically. Preserve names, numbers, code, formulas, URLs, and domain terminology. "
                        "Never answer the source content. IF THE TARGET LANGUAGE IS TELUGU ('te'), STRICTLY RETURN ONLY TELUGU. ABSOLUTELY DO NOT RETURN HINDI UNDER ANY CIRCUMSTANCES. In the 'translation' field, return ONLY Telugu script, no transliteration, no parentheses, and no English explanations. If the source content is grammatically ambiguous or incomplete, infer the most likely natural human intent rather than translating mechanically or literally. Return only valid JSON with keys: translation, "
                        "detected_source_language, transliteration, alternatives, notes, confidence, sentence_breakdown. "
                        "confidence must be a number from 0 to 1; alternatives and notes must be arrays of strings. "
                        "For a short sentence or phrase, sentence_breakdown must be an object with keys: original, "
                        "translation, word_by_word, overall_meaning, formal_version, informal_version, natural_version, "
                        "grammar_notes. word_by_word is an array of {source, meaning, translation, transliteration, note}. "
                        "Use the target language script in translation/versions, include romanization in transliteration or "
                        "parentheses where useful, and explain grammar briefly in English. For long multi-sentence text, "
                        "set sentence_breakdown to null."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Source: {source_label}\nTarget: {target_label} ({target})\n"
                        f"Tone: {request.tone}\n{formatting}\n{alternatives}{context}{section}\n\n"
                        f"Content:\n{content}"
                    ),
                },
            ],
            max_output_tokens=max(700, min(3500, len(content) // 2 + 600)),
        )

    async def translate_text(
        self,
        content: str,
        source_lang: LanguageCode,
        target_lang: LanguageCode,
        tone: ToneStyle | None = None,
        preserve_formatting: bool = True,
    ) -> TranslationTextResponse:
        return await self.translate(
            TranslationTextRequest(
                content=content,
                source_language=source_lang.value,
                target_language=target_lang.value,
                tone=(tone or ToneStyle.PRESERVE).value,
                preserve_formatting=preserve_formatting,
            )
        )

    async def grammar(self, request: GrammarRequest) -> GrammarResponse:
        language = self.validate_language(request.language, allow_auto=True)
        payload = await self._json_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are a multilingual writing editor. Preserve meaning, facts, names, and formatting. "
                        "Return only JSON with corrected_content, detected_language, and explanation (array of short strings)."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Language: {language}. Editing goal: {request.goal}.\n"
                        f"Edit only what improves the writing.\n\n{request.content}"
                    ),
                },
            ]
        )
        corrected = str(payload.get("corrected_content", "")).strip()
        if not corrected:
            raise HTTPException(status_code=502, detail="The writing model returned an empty response")
        return GrammarResponse(
            original_content=request.content,
            corrected_content=corrected,
            language=self._normalize_detected_language(str(payload.get("detected_language", language))),
            explanation=self._string_list(payload.get("explanation")),
        )

    async def analyze(self, request: AnalysisRequest) -> AnalysisResponse:
        source = self.validate_language(request.source_language, allow_auto=True)
        target = self.validate_language(request.target_language)
        if request.analysis_type == "word":
            keys = (
                "original, pronunciation, ipa, meaning, translation, part_of_speech, root_word, "
                "synonyms, antonyms, examples, usage"
            )
        else:
            keys = (
                "original, translation, word_by_word (array of objects with source, meaning, translation, "
                "transliteration, note), grammar, grammar_notes, meaning, overall_meaning, context, alternatives, "
                "formal_version, informal_version, natural_version, professional_version, simple_version, spoken_version"
            )
        payload = await self._json_chat(
            [
                {
                    "role": "system",
                    "content": (
                        "You are an expert multilingual linguist and teacher. Return only valid JSON. "
                        f"Top-level keys: detected_source_language and analysis. Analysis keys: {keys}. "
                        "Use arrays for synonyms, antonyms, examples, grammar, grammar_notes, alternatives, and word_by_word. "
                        "For sentence analysis, follow this teaching structure: Word-for-Word Meanings, Overall Sentence "
                        "Meaning, Formal/Polite version, Informal version, and short grammar notes."
                    ),
                },
                {
                    "role": "user",
                    "content": (
                        f"Analyze this {request.analysis_type}. Source language: {source}. "
                        f"Translate explanations and the main translation into {self.language_name(target)} ({target}).\n\n"
                        f"{request.content}"
                    ),
                },
            ]
        )
        data = payload.get("analysis")
        if not isinstance(data, dict):
            raise HTTPException(status_code=502, detail="The language model returned invalid analysis")
        data.setdefault("original", request.content)
        detected = self._normalize_detected_language(str(payload.get("detected_source_language", source)))
        if request.analysis_type == "word":
            return AnalysisResponse(analysis_type="word", detected_source_language=detected, word=WordAnalysis(**data))
        data = self._normalize_sentence_data(data, original=request.content, translation=str(data.get("translation", "")))
        return AnalysisResponse(
            analysis_type="sentence",
            detected_source_language=detected,
            sentence=SentenceAnalysis(**data),
        )

    async def extract_upload(self, file: UploadFile) -> tuple[str, str]:
        content = await file.read()
        if not content:
            raise HTTPException(status_code=400, detail="The uploaded file is empty")
        if len(content) > 15 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Files must be 15 MB or smaller")

        filename = file.filename or "upload"
        suffix = Path(filename).suffix.lower()
        content_type = file.content_type or "application/octet-stream"
        try:
            if content_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}:
                try:
                    Image.open(BytesIO(content)).verify()
                except UnidentifiedImageError as exc:
                    raise HTTPException(status_code=400, detail="Upload a valid image file") from exc
                extracted = VisionService().ocr_bytes(content).strip()
                return self._validate_extracted_text(extracted), "image"
            if suffix == ".pdf" or content_type == "application/pdf":
                reader = PdfReader(BytesIO(content))
                extracted = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
                return self._validate_extracted_text(extracted), "document"
            if suffix == ".docx":
                document = WordDocument(BytesIO(content))
                extracted = "\n".join(paragraph.text for paragraph in document.paragraphs).strip()
                return self._validate_extracted_text(extracted), "document"
            if suffix in {".txt", ".md", ".csv", ".json", ".html"} or content_type.startswith("text/"):
                extracted = content.decode("utf-8", errors="ignore").strip()
                return self._validate_extracted_text(extracted), "document"
        except TesseractNotFoundError as exc:
            raise HTTPException(status_code=503, detail="Tesseract OCR is not installed or available on PATH") from exc
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail="The uploaded file could not be read") from exc
        raise HTTPException(status_code=415, detail="Use an image, PDF, DOCX, TXT, Markdown, CSV, JSON, or HTML file")

    async def save_history(
        self,
        user_id: str,
        original: str,
        result: TranslationTextResponse,
        mode: str = "text",
    ) -> None:
        if not self.session:
            return
        try:
            self.session.add(
                TranslationHistory(
                    user_id=user_id,
                    source_language=result.detected_source_language,
                    target_language=result.target_language,
                    original_content=original,
                    translated_content=result.translated_content,
                    mode=mode,
                    character_count=len(original),
                )
            )
            await self.session.commit()
        except Exception as exc:
            logger.warning("Failed to save translation history: %s", exc)
            with suppress(Exception):
                await self.session.rollback()

    async def history(self, user_id: str, limit: int = 20) -> list[TranslationHistory]:
        if not self.session:
            return []
        result = await self.session.execute(
            select(TranslationHistory)
            .where(TranslationHistory.user_id == user_id)
            .order_by(desc(TranslationHistory.created_at))
            .limit(limit)
        )
        return list(result.scalars().all())

    async def clear_history(self, user_id: str) -> None:
        if not self.session:
            return
        rows = await self.history(user_id, 1000)
        for row in rows:
            await self.session.delete(row)
        await self.session.commit()

    async def ocr_with_boxes(self, content: bytes, *, language: str = "auto") -> OCRResponse:
        try:
            image = Image.open(BytesIO(content))
        except UnidentifiedImageError as exc:
            raise HTTPException(status_code=400, detail="Upload a valid image file") from exc

        try:
            data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
        except TesseractNotFoundError as exc:
            raise HTTPException(status_code=503, detail="Tesseract OCR is not installed or is not available on PATH") from exc

        boxes: list[OCRBox] = []
        texts: list[str] = []
        n = len(data.get("text", []))
        for i in range(n):
            text = (data["text"][i] or "").strip()
            if not text:
                continue
            conf = int(data["conf"][i]) if data["conf"][i] not in (None, "") else 0
            boxes.append(
                OCRBox(
                    text=text,
                    x=int(data["left"][i]),
                    y=int(data["top"][i]),
                    width=int(data["width"][i]),
                    height=int(data["height"][i]),
                    confidence=max(0, conf),
                )
            )
            texts.append(text)

        full_text = " ".join(texts).strip()
        if not full_text:
            return OCRResponse(text="", boxes=[], language=language, mode="image")

        return OCRResponse(text=full_text, boxes=boxes, language=language, mode="image")

    async def explain(self, payload: ExplainRequest) -> ExplainResponse:
        source = self.validate_language(payload.source_language, allow_auto=True)
        target = self.validate_language(payload.target_language)
        focus_instruction = {
            "grammar": "Focus primarily on grammar rules, tenses, and sentence structure.",
            "vocabulary": "Focus primarily on word meanings, synonyms, and usage.",
            "usage": "Focus primarily on real-life usage, context, and common mistakes.",
            "culture": "Focus primarily on cultural context, idioms, and nuances.",
            "all": "Provide a comprehensive explanation covering grammar, vocabulary, usage, and cultural context.",
        }.get(payload.focus, "Provide a comprehensive explanation.")

        prompt = (
            "You are an expert language teacher and translator. Explain the translation between the source and target languages. "
            "Return only valid JSON with keys: original, translation, why_correct (array), grammar_rules (array), "
            "vocabulary_notes (array), real_life_usage (array), common_mistakes (array), better_alternatives (array). "
            f"Source language: {self.language_name(source)}. Target language: {self.language_name(target)}. "
            f"{focus_instruction}\n\nContent:\n{payload.content}"
        )
        data = await self._json_chat([
            {"role": "system", "content": "You are an expert language teacher and translator. Return only valid JSON."},
            {"role": "user", "content": prompt},
        ], max_output_tokens=2000)

        return ExplainResponse(
            original=str(data.get("original", payload.content)).strip(),
            translation=str(data.get("translation", "")).strip(),
            why_correct=self._string_list(data.get("why_correct")),
            grammar_rules=self._string_list(data.get("grammar_rules")),
            vocabulary_notes=self._string_list(data.get("vocabulary_notes")),
            real_life_usage=self._string_list(data.get("real_life_usage")),
            common_mistakes=self._string_list(data.get("common_mistakes")),
            better_alternatives=self._string_list(data.get("better_alternatives")),
        )

    async def generate_flashcards(self, payload: FlashcardGenerateRequest) -> list[FlashcardResponse]:
        source = self.validate_language(payload.source_language, allow_auto=True)
        target = self.validate_language(payload.target_language)
        prompt = (
            f"Generate {payload.count} flashcards for learning {self.language_name(target)} from {self.language_name(source)}. "
            "Each flashcard should have a front (word/phrase in source language) and back (translation + brief meaning in target language). "
            "Return only valid JSON with key 'flashcards' as an array of objects with keys: front, back. "
            f"Content to base flashcards on:\n{payload.content}"
        )
        data = await self._json_chat([
            {"role": "system", "content": "You are an expert language teacher. Create effective flashcards. Return only valid JSON."},
            {"role": "user", "content": prompt},
        ], max_output_tokens=2000)

        raw_cards = data.get("flashcards") or []
        return [
            FlashcardResponse(
                id="",
                front=str(card.get("front", "")).strip(),
                back=str(card.get("back", "")).strip(),
                source_language=source,
                target_language=target,
            )
            for card in raw_cards
            if isinstance(card, dict) and str(card.get("front", "")).strip()
        ][: payload.count]

    async def generate_quiz(self, payload: QuizGenerateRequest) -> QuizResponse:
        source = self.validate_language(payload.source_language, allow_auto=True)
        target = self.validate_language(payload.target_language)
        types_str = ", ".join(payload.types)
        prompt = (
            f"Generate a {payload.count}-question language learning quiz from {self.language_name(source)} to {self.language_name(target)}. "
            f"Question types: {types_str}. Each question should test vocabulary, grammar, or translation. "
            "Return only valid JSON with key 'questions' as an array of objects with keys: question, options (array), answer, explanation. "
            f"Content to base quiz on:\n{payload.content}"
        )
        data = await self._json_chat([
            {"role": "system", "content": "You are an expert language teacher. Create engaging quizzes. Return only valid JSON."},
            {"role": "user", "content": prompt},
        ], max_output_tokens=3000)

        raw_questions = data.get("questions") or []
        questions = [
            QuizQuestionSchema(
                question=str(q.get("question", "")).strip(),
                options=[str(o).strip() for o in (q.get("options") or []) if str(o).strip()],
                answer=str(q.get("answer", "")).strip(),
                explanation=str(q.get("explanation", "")).strip(),
            )
            for q in raw_questions
            if isinstance(q, dict) and str(q.get("question", "")).strip()
        ][: payload.count]

        return QuizResponse(
            id="",
            topic="Generated Quiz",
            source_language=source,
            target_language=target,
            questions=questions,
            total=len(questions),
        )

    async def vocabulary_list(self, user_id: str, db: AsyncSession | None = None) -> list[VocabularyItem]:
        if not db:
            return []
        result = await db.execute(
            select(VocabularyItem).where(VocabularyItem.user_id == user_id).order_by(desc(VocabularyItem.created_at))
        )
        return list(result.scalars().all())

    async def add_vocabulary(self, user_id: str, item: VocabularyItem, db: AsyncSession | None = None) -> VocabularyItem:
        if not db:
            return item
        db.add(item)
        await db.commit()
        await db.refresh(item)
        return item

    async def remove_vocabulary(self, item_id: str, db: AsyncSession | None = None) -> None:
        if not db:
            return
        item = await db.get(VocabularyItem, item_id)
        if item:
            await db.delete(item)
            await db.commit()

    async def bookmark_list(self, user_id: str, db: AsyncSession | None = None) -> list[Bookmark]:
        if not db:
            return []
        result = await db.execute(
            select(Bookmark).where(Bookmark.user_id == user_id).order_by(desc(Bookmark.created_at))
        )
        return list(result.scalars().all())

    async def add_bookmark(self, user_id: str, bookmark: Bookmark, db: AsyncSession | None = None) -> Bookmark:
        if not db:
            return bookmark
        db.add(bookmark)
        await db.commit()
        await db.refresh(bookmark)
        return bookmark

    async def remove_bookmark(self, bookmark_id: str, db: AsyncSession | None = None) -> None:
        if not db:
            return
        item = await db.get(Bookmark, bookmark_id)
        if item:
            await db.delete(item)
            await db.commit()

    async def saved_phrase_list(self, user_id: str, db: AsyncSession | None = None) -> list[SavedPhrase]:
        if not db:
            return []
        result = await db.execute(
            select(SavedPhrase).where(SavedPhrase.user_id == user_id).order_by(desc(SavedPhrase.created_at))
        )
        return list(result.scalars().all())

    async def add_saved_phrase(self, user_id: str, phrase: SavedPhrase, db: AsyncSession | None = None) -> SavedPhrase:
        if not db:
            return phrase
        db.add(phrase)
        await db.commit()
        await db.refresh(phrase)
        return phrase

    async def remove_saved_phrase(self, phrase_id: str, db: AsyncSession | None = None) -> None:
        if not db:
            return
        item = await db.get(SavedPhrase, phrase_id)
        if item:
            await db.delete(item)
            await db.commit()

    async def save_quiz(self, user_id: str, quiz: Quiz, db: AsyncSession | None = None) -> Quiz:
        if not db:
            return quiz
        db.add(quiz)
        await db.commit()
        await db.refresh(quiz)
        return quiz

    async def list_quizzes(self, user_id: str, db: AsyncSession | None = None) -> list[Quiz]:
        if not db:
            return []
        result = await db.execute(
            select(Quiz).where(Quiz.user_id == user_id).order_by(desc(Quiz.created_at)).limit(50)
        )
        return list(result.scalars().all())

    async def _json_chat(
        self,
        messages: list[dict[str, str]],
        *,
        max_output_tokens: int | None = None,
    ) -> dict[str, Any]:
        # Deterministic model calls are cacheable by their prompt content.
        # This makes repeated identical translations/grammar/analysis instant.
        import hashlib

        cache_key = "tjson:" + hashlib.sha256(
            json.dumps(messages, ensure_ascii=False, sort_keys=True).encode("utf-8")
        ).hexdigest()
        cached = await translation_cache.get(cache_key)
        if cached is not None:
            return cached

        last_exc: Exception | None = None
        for attempt in range(4):
            try:
                raw = await self.gateway.chat(
                    messages,
                    model=self.gateway.settings.ollama_translation_model,
                    temperature=0.1,
                    max_output_tokens=max_output_tokens,
                    json_mode=True,
                    context_window=self.gateway.settings.translation_context_window,
                    request_timeout_seconds=self.gateway.settings.translation_request_timeout_seconds,
                )
                parsed = _clean_json(raw)
                if not isinstance(parsed, dict) or not parsed:
                    raise ValueError("The translation model returned an empty JSON object")
                await translation_cache.set(cache_key, parsed, ttl_seconds=3600)
                return parsed
            except HTTPException:
                raise
            except (httpx.HTTPError, TimeoutError, RuntimeError) as exc:
                last_exc = exc
                logger.warning("Translation model unavailable (attempt %d): %s", attempt + 1, exc)
                if attempt == 3:
                    break
                sleep = min(2 ** attempt + random.uniform(0, 1), 12)
                await asyncio.sleep(sleep)
            except (ValueError, json.JSONDecodeError) as exc:
                last_exc = exc
                logger.warning("Invalid structured translation response (attempt %d): %s", attempt + 1, exc)
                if attempt == 3:
                    break
                sleep = min(2 ** attempt + random.uniform(0, 1), 8)
                await asyncio.sleep(sleep)

        if isinstance(last_exc, (httpx.HTTPError, TimeoutError, RuntimeError)):
            raise HTTPException(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                detail="The translation model is unavailable. Check Gemini/Ollama settings and try again.",
            ) from last_exc

        try:
            fallback_raw = await self.gateway.chat(
                messages,
                model=self.gateway.settings.ollama_translation_model,
                temperature=0.1,
                max_output_tokens=max_output_tokens,
                json_mode=False,
                context_window=self.gateway.settings.translation_context_window,
                request_timeout_seconds=self.gateway.settings.translation_request_timeout_seconds,
            )
            fallback_text = str(fallback_raw).strip()
            if fallback_text:
                detected = "unknown"
                for message in messages:
                    if "Source:" in message.get("content", ""):
                        detected = message["content"].split("Source:")[-1].split("\n")[0].strip()
                        break
                parsed = {
                    "translation": fallback_text,
                    "detected_source_language": detected,
                    "transliteration": "",
                    "alternatives": [],
                    "notes": ["Translated without structured JSON mode."],
                    "confidence": 0.8,
                }
                await translation_cache.set(cache_key, parsed, ttl_seconds=3600)
                return parsed
        except Exception:
            pass

        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="The model could not produce a valid structured translation. Try again.",
        ) from last_exc

    @staticmethod
    def _normalize_sentence_data(data: dict[str, Any], *, original: str, translation: str) -> dict[str, Any]:
        normalized = dict(data)
        normalized.setdefault("original", original)
        if translation and not str(normalized.get("translation", "")).strip():
            normalized["translation"] = translation

        aliases = {
            "overall_meaning": ("overall_sentence_meaning", "sentence_meaning", "meaning"),
            "formal_version": ("formal", "formal_polite", "formal_polite_version", "polite_version"),
            "informal_version": ("informal", "casual", "casual_version"),
            "natural_version": ("natural", "natural_translation", "spoken_version"),
        }
        for target_key, source_keys in aliases.items():
            if str(normalized.get(target_key, "")).strip():
                continue
            for source_key in source_keys:
                value = normalized.get(source_key)
                if isinstance(value, str) and value.strip():
                    normalized[target_key] = value.strip()
                    break

        word_items = (
            normalized.get("word_by_word")
            or normalized.get("word_for_word")
            or normalized.get("word_for_word_meanings")
            or []
        )
        normalized_words: list[dict[str, str]] = []
        if isinstance(word_items, list):
            for item in word_items:
                if isinstance(item, dict):
                    source = str(
                        item.get("source")
                        or item.get("word")
                        or item.get("original")
                        or item.get("term")
                        or ""
                    ).strip()
                    meaning = str(
                        item.get("meaning")
                        or item.get("definition")
                        or item.get("literal_meaning")
                        or item.get("target_meaning")
                        or ""
                    ).strip()
                    translated = str(
                        item.get("translation")
                        or item.get("target")
                        or item.get("translated")
                        or item.get("target_word")
                        or ""
                    ).strip()
                    transliteration = str(
                        item.get("transliteration")
                        or item.get("romanization")
                        or item.get("pronunciation")
                        or ""
                    ).strip()
                    note = str(item.get("note") or item.get("usage") or "").strip()
                else:
                    source = str(item).strip()
                    meaning = ""
                    translated = ""
                    transliteration = ""
                    note = ""
                if source:
                    normalized_words.append(
                        {
                            "source": source,
                            "meaning": meaning,
                            "translation": translated,
                            "transliteration": transliteration,
                            "note": note,
                        }
                    )
        normalized["word_by_word"] = normalized_words

        grammar_notes = normalized.get("grammar_notes") or normalized.get("grammar") or []
        normalized["grammar_notes"] = TranslationEngine._string_list(grammar_notes)
        if not normalized.get("grammar"):
            normalized["grammar"] = normalized["grammar_notes"]
        if not str(normalized.get("meaning", "")).strip():
            normalized["meaning"] = str(normalized.get("overall_meaning", "")).strip()
        return normalized

    @staticmethod
    def _sentence_breakdown(value: Any, *, original: str, translation: str) -> SentenceBreakdown | None:
        if not isinstance(value, dict):
            return None
        data = TranslationEngine._normalize_sentence_data(value, original=original, translation=translation)
        has_teaching_detail = bool(
            data.get("word_by_word")
            or str(data.get("overall_meaning", "")).strip()
            or str(data.get("formal_version", "")).strip()
            or str(data.get("informal_version", "")).strip()
            or str(data.get("natural_version", "")).strip()
            or data.get("grammar_notes")
        )
        if not has_teaching_detail:
            return None
        try:
            return SentenceBreakdown(**data)
        except ValidationError as exc:
            logger.warning("Invalid sentence breakdown: %s", exc)
            return None

    @staticmethod
    def _string_list(value: Any) -> list[str]:
        if not isinstance(value, list):
            return []
        return [str(item).strip() for item in value if str(item).strip()][:8]

    @staticmethod
    def _confidence(value: Any) -> float:
        try:
            return max(0.0, min(1.0, float(value)))
        except (TypeError, ValueError):
            return 0.85

    @staticmethod
    def _normalize_detected_language(value: str) -> str:
        normalized = value.strip()
        match = LANGUAGE_BY_CODE.get(normalized.lower())
        if match:
            return match[0]
        name_match = LANGUAGE_BY_NAME.get(normalized.lower())
        if name_match:
            return name_match
        base_code = normalized.split("-", 1)[0].lower()
        base_match = LANGUAGE_BY_CODE.get(base_code)
        return base_match[0] if base_match else normalized or "unknown"

    @staticmethod
    def language_name(code: str) -> str:
        match = LANGUAGE_BY_CODE.get(code.lower())
        return match[1] if match else code

    @staticmethod
    def _validate_extracted_text(content: str) -> str:
        if len(content) > 50000:
            raise HTTPException(status_code=413, detail="The extracted text exceeds the 50,000 character limit")
        return content

    @staticmethod
    def _split_content(content: str, max_chars: int = 10000) -> list[str]:
        if len(content) <= max_chars:
            return [content]
        chunks: list[str] = []
        current: list[str] = []
        current_size = 0
        for paragraph in content.splitlines():
            line = paragraph.strip()
            if not line:
                if current:
                    current.append("")
                continue
            while len(line) > max_chars:
                if current:
                    chunks.append("\n".join(current).strip())
                    current = []
                    current_size = 0
                split_at = line.rfind(" ", 0, max_chars)
                split_at = split_at if split_at > max_chars // 2 else max_chars
                chunks.append(line[:split_at].strip())
                line = line[split_at:].strip()
            projected = current_size + len(line) + 1
            if current and projected > max_chars:
                chunks.append("\n".join(current).strip())
                current = [line]
                current_size = len(line)
            else:
                current.append(line)
                current_size = projected
        if current:
            chunks.append("\n".join(current).strip())
        return [chunk for chunk in chunks if chunk]


__all__ = [
    "ContentType",
    "LanguageCode",
    "LanguagePair",
    "SavedTranslations",
    "ToneStyle",
    "TranslationEngine",
    "TranslationHistory",
    "TranslationMode",
    "TranslationRequest",
    "TranslationResult",
    "VoiceTranslationResult",
    "supported_languages",
]
